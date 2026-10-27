"""Reproduce the four version-one repository defaults; never runtime uploads.

Run with python-docx. Existing released defaults must not be changed in place:
future layouts use a new filename/version and retain the old default.
"""
from pathlib import Path
from io import BytesIO
import sys
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.services.api.disposition.template_contract import CONTRACTS


def prop(parent,name,**values):
    node=OxmlElement('w:'+name)
    for key,value in values.items():node.set(qn('w:'+key),str(value))
    parent.append(node);return node


def paragraph(host,value='',*,rtl=False,style=None,size=None):
    p=host.add_paragraph(style=style)
    prop(p._p.get_or_add_pPr(),'bidi',val=int(rtl))
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:p._p.get_or_add_pPr().jc.set(qn('w:val'),'start')
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=style in ('Heading 1','Heading 2')
    r=p.add_run(value);r.font.name='Changa';r.font.color.rgb=RGBColor(0,0,0)
    fonts=r._r.get_or_add_rPr().rFonts
    fonts.set(qn('w:cs'),'Changa');fonts.set(qn('w:eastAsia'),'Changa')
    if rtl:prop(r._r.get_or_add_rPr(),'rtl',val=1)
    if size:r.font.size=Pt(size);prop(r._r.get_or_add_rPr(),'szCs',val=int(size*2))
    return p


def field(host,key,label):
    paragraph(host,label[0],rtl=True,size=9)
    paragraph(host,label[1],size=8)
    paragraph(host,'{{'+key+'}}',size=9)


def table(doc,widths):
    t=doc.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for column,width in zip(t.columns,widths):column.width=Mm(width)
    pr=t._tbl.tblPr;prop(pr,'bidiVisual',val=1)
    borders=prop(pr,'tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):prop(borders,side,val='single',sz=4,color='D9D9D9')
    margins=prop(pr,'tblCellMar')
    for side in ('top','bottom','left','right'):prop(margins,side,w=90,type='dxa')
    return t


def main():
    target=ROOT/'templates/disposition/destruction';target.mkdir(parents=True,exist_ok=True)
    for kind,c in CONTRACTS.items():
        # Working layouts have dedicated builders; never overwrite them here.
        if not c.filename.endswith('.default.v1.docx'):continue
        doc=Document();section=doc.sections[0]
        section.page_width=Mm(210);section.page_height=Mm(297)
        section.top_margin=section.bottom_margin=Mm(17);section.left_margin=section.right_margin=Mm(18)
        for style in doc.styles:
            if style.type in (1,2):
                style.font.name='Changa';style.font.size=Pt(10);style.font.color.rgb=RGBColor(0,0,0)
                fonts=style.element.get_or_add_rPr().rFonts
                fonts.set(qn('w:cs'),'Changa');fonts.set(qn('w:eastAsia'),'Changa')
                style.font.italic=False
                for border in style.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
        paragraph(doc,c.arabic,rtl=True,style='Title',size=21)
        paragraph(doc,c.english,style='Subtitle',size=15)
        # Two columns keep metadata readable; cells flow without fixed heights.
        t=table(doc,[87,87]);items=list(c.fields.items())
        for index in range(0,len(items),2):
            row=t.add_row();prop(row._tr.get_or_add_trPr(),'cantSplit',val=1);cells=row.cells
            for cell,(key,label) in zip(cells,items[index:index+2]):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell._tc.remove(cell.paragraphs[0]._p);field(cell,key,label)
        if c.group:
            paragraph(doc,'التفاصيل',rtl=True,style='Heading 1',size=13)
            paragraph(doc,'Details',style='Heading 2',size=11)
            t=table(doc,[174/len(c.columns)]*len(c.columns))
            header=t.add_row();prop(header._tr.get_or_add_trPr(),'tblHeader',val=1)
            labels={
              'destruction_request':[('التصنيف ومصدر القاعدة','Classification and rule source'),('قاعدة الحفظ الفعلية','Effective retention rule'),('أعداد العناصر','Item counts'),('وسيط الحفظ','Record medium')],
              'proposal_execution_report':[('المقترح وصاحبه','Proposal and proposer'),('نتيجة التنفيذ','Execution result'),('الاستبعادات وسجل التغييرات','Removed units and history')],
              'physical_preparation_report':[('التصنيف والملف','Classification and aggregation'),('الملكية والمواقع','Ownership and locations'),('أعداد العناصر التابعة','Descendant counts')],
            }[kind]
            for cell,labels in zip(header.cells,labels):
                cell._tc.remove(cell.paragraphs[0]._p)
                prop(cell._tc.get_or_add_tcPr(),'shd',fill='EEF7FD')
                paragraph(cell,labels[0],rtl=True,size=10).paragraph_format.keep_with_next=True
                paragraph(cell,labels[1],size=9).paragraph_format.keep_with_next=True
            row=t.add_row();prop(row._tr.get_or_add_trPr(),'cantSplit',val=1);cells=row.cells
            for i,(cell,column) in enumerate(zip(cells,c.columns)):
                cell._tc.remove(cell.paragraphs[0]._p)
                if i==0:paragraph(cell,'{{#'+c.group+'}}',size=1)
                for key,label in column.items():field(cell,key,label)
                if i==len(c.columns)-1:paragraph(cell,'{{/'+c.group+'}}',size=1)
        doc.core_properties.title=c.english;doc.core_properties.subject=c.arabic
        doc.core_properties.author='Wathiq';doc.core_properties.comments=''
        # Page numbers are dynamic Word fields, not application placeholders.
        p=paragraph(section.footer,'',size=8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field_node=OxmlElement('w:fldSimple');field_node.set(qn('w:instr'),'PAGE');p._p.append(field_node)
        package=BytesIO();doc.save(package)
        with ZipFile(package) as source,ZipFile(target/c.filename,'w',ZIP_DEFLATED) as output:
            for name in source.namelist():
                info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED
                output.writestr(info,source.read(name))
        print(c.filename)


if __name__=='__main__':main()
