"""Author the working execution report without overwriting the retained v1."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import sys
from docx import Document
from docx.shared import Mm
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from tools.build_disposition_templates import paragraph,table,prop
from tools.build_request_template_v2 import request_field,organization_header
from backend.services.api.disposition.template_contract import CONTRACTS,group_fields


def repeat(doc,group,columns):
    t=table(doc,[width for _,_,width in columns]);header=t.add_row();prop(header._tr.get_or_add_trPr(),'tblHeader',val=1)
    for cell,(_,label,_) in zip(header.cells,columns):
        cell._tc.remove(cell.paragraphs[0]._p);prop(cell._tc.get_or_add_tcPr(),'shd',fill='EEF7FD')
        paragraph(cell,label[0],rtl=True,size=9);paragraph(cell,label[1],size=8)
    row=t.add_row()
    for i,(cell,(keys,_,_)) in enumerate(zip(row.cells,columns)):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell._tc.remove(cell.paragraphs[0]._p)
        if i==0:paragraph(cell,'{{#'+group+'}}',size=1)
        for key in keys:
            if len(keys)>1:
                label=group_fields(CONTRACTS['proposal_execution_report'])[group][key]
                paragraph(cell,label[0]+' / '+label[1],rtl=True,size=7)
            paragraph(cell,'{{'+key+'}}',size=8)
        if i==len(columns)-1:paragraph(cell,'{{/'+group+'}}',size=1)


def main():
    directory=ROOT/'templates/disposition/destruction';c=CONTRACTS['proposal_execution_report']
    doc=Document(directory/'proposal-execution-report.default.v1.docx')
    for node in list(doc.element.body):
        if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
    organization_header(doc,c.fields['organization.name'])
    paragraph(doc,c.arabic,rtl=True,style='Title',size=20);paragraph(doc,c.english,style='Subtitle',size=13)
    meta=table(doc,[87,87]);items=[(k,v) for k,v in c.fields.items() if k not in c.optional and k!='organization.name']
    for i in range(0,len(items),2):
        row=meta.add_row();prop(row._tr.get_or_add_trPr(),'cantSplit',val=1)
        for cell,(key,label) in zip(row.cells,items[i:i+2]):
            cell._tc.remove(cell.paragraphs[0]._p);request_field(cell,key,label)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE);section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Mm(297);section.page_height=Mm(210)
    section.top_margin=section.bottom_margin=Mm(17);section.left_margin=section.right_margin=Mm(18)
    paragraph(doc,'الملفات المتأثرة بالمقترحات المنفّذة',rtl=True,style='Heading 1',size=15)
    paragraph(doc,'Aggregations affected by executed proposals',style='Heading 2',size=11)
    labels=group_fields(c)['aggregations']
    widths=[20,42,55,18,18,29,22,57]
    repeat(doc,'aggregations',[( (key,),label,width) for (key,label),width in zip(labels.items(),widths)])
    paragraph(doc,'المقترحات المنفّذة',rtl=True,style='Heading 1',size=15).paragraph_format.page_break_before=True
    paragraph(doc,'Executed proposals',style='Heading 2',size=11)
    repeat(doc,'proposals',[
        (('proposal.serial_number',),('الرقم التسلسلي','No.'),20),
        (('proposal.action','proposal.target'),('الإجراء والعنصر المستهدف','Action and target'),53),
        (('proposal.change',),('التغيير المنفّذ','Executed change'),76),
        (('proposal.proposer','proposal.executor'),('مقدّم المقترح والمنفّذ','Proposer and executor'),54),
        (('proposal.result','proposal.reason','proposal.removed_units'),('نتيجة التنفيذ','Execution outcome'),58),
    ])
    doc.core_properties.title=c.english;doc.core_properties.subject=c.arabic
    package=BytesIO();doc.save(package)
    with ZipFile(package) as source,ZipFile(directory/c.filename,'w',ZIP_DEFLATED) as out:
        for name in source.namelist():
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED
            out.writestr(info,source.read(name))
    print(directory/c.filename)

if __name__=='__main__':main()
