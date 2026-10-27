"""Author request v2 from the retained v1 package; never overwrite a release."""
from io import BytesIO
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
from pathlib import Path
import sys
from docx import Document
from docx.shared import Pt,Mm
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from tools.build_disposition_templates import paragraph,field,table,prop
from backend.services.api.disposition.template_contract import CONTRACTS


def request_field(host,key,label):
    paragraph(host,label[0]+' / '+label[1],rtl=True,size=8)
    paragraph(host,'{{'+key+'}}',size=9)


def organization_header(host,label):
    paragraph(host,label[0]+' / '+label[1],rtl=True,size=8)
    name=paragraph(host,'{{organization.name}}',size=16)
    name.paragraph_format.keep_with_next=True
    name.paragraph_format.keep_together=True
    name.paragraph_format.space_after=Pt(6)
    for run in name.runs:
        run.bold=True
        prop(run._r.get_or_add_rPr(),'bCs',val=1)


def main():
    directory=ROOT/'templates/disposition/destruction'
    doc=Document(directory/'request-for-destruction.default.v1.docx')
    # Keep the released styles, section geometry, footer and package resources.
    for node in list(doc.element.body):
        if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
    c=CONTRACTS['destruction_request']
    organization_header(doc,c.fields['organization.name'])
    contacts=table(doc,[87,87]);row=contacts.add_row()
    for cell,key in zip(row.cells,('organization.email','organization.phone')):
        cell._tc.remove(cell.paragraphs[0]._p);request_field(cell,key,c.fields[key])
    paragraph(doc,c.arabic,rtl=True,style='Title',size=20)
    paragraph(doc,c.english,style='Subtitle',size=13)
    items=[(key,label) for key,label in c.fields.items() if key not in c.optional and not key.startswith('organization.')]
    metadata=table(doc,[87,87])
    for index in range(0,len(items),2):
        row=metadata.add_row();prop(row._tr.get_or_add_trPr(),'cantSplit',val=1)
        for cell,(key,label) in zip(row.cells,items[index:index+2]):
            cell._tc.remove(cell.paragraphs[0]._p);request_field(cell,key,label)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation=WD_ORIENT.LANDSCAPE;section.page_width=Mm(297);section.page_height=Mm(210)
    section.top_margin=section.bottom_margin=Mm(17);section.left_margin=section.right_margin=Mm(18)
    paragraph(doc,'الملفات المطلوب إتلافها',rtl=True,style='Heading 1',size=15)
    paragraph(doc,'Aggregations requested for destruction',style='Heading 2',size=11)
    columns=(
      ('aggregation.serial_number',('م','No.'),12),
      ('aggregation.number',('رقم الملف','Aggregation number'),42),
      ('aggregation.title',('عنوان الملف','Aggregation title'),59),
      ('aggregation.from_year',('من سنة','From Year'),18),
      ('aggregation.to_year',('إلى سنة','To Year'),18),
      ('aggregation.classification_code',('رقم التصنيف','Classification number'),29),
      ('aggregation.records',('إجمالي الوثائق','Total records'),22),
      ('aggregation.assigned_location',('الموقع المعيّن للورقي أو المختلط','Assigned location physical or mixed'),61),
    )
    t=table(doc,[column[2] for column in columns]);header=t.add_row();prop(header._tr.get_or_add_trPr(),'tblHeader',val=1)
    for cell,(_,label,_) in zip(header.cells,columns):
        cell._tc.remove(cell.paragraphs[0]._p);prop(cell._tc.get_or_add_tcPr(),'shd',fill='EEF7FD')
        paragraph(cell,label[0],rtl=True,size=9);paragraph(cell,label[1],size=8)
    row=t.add_row()
    for index,(cell,(key,_,_)) in enumerate(zip(row.cells,columns)):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell._tc.remove(cell.paragraphs[0]._p)
        if index==0:paragraph(cell,'{{#aggregations}}',size=1)
        paragraph(cell,'{{'+key+'}}',size=9)
        if index==len(columns)-1:paragraph(cell,'{{/aggregations}}',size=1)
    doc.core_properties.title=c.english;doc.core_properties.subject=c.arabic
    package=BytesIO();doc.save(package)
    with ZipFile(package) as source,ZipFile(directory/c.filename,'w',ZIP_DEFLATED) as output:
        for name in source.namelist():
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED
            output.writestr(info,source.read(name))
    print(directory/c.filename)

if __name__=='__main__':main()
