"""Author the concise working certificate v2; preserve the released v1 package."""
from io import BytesIO
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
from pathlib import Path
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from tools.build_disposition_templates import paragraph,table,prop
from backend.services.api.disposition.template_contract import CONTRACTS


def main():
    directory=ROOT/'templates/disposition/destruction'
    doc=Document(directory/'digital-destruction-certificate.default.v1.docx')
    for node in list(doc.element.body):
        if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
    c=CONTRACTS['digital_destruction_certificate']
    # Arabic opening label satisfies bilingual template validation; the prominent
    # organization value follows without imposing a language on a proper name.
    paragraph(doc,'اسم الجهة / Organization',rtl=True,size=8)
    p=paragraph(doc,'{{organization.name}}',size=17)
    p.paragraph_format.space_after=Pt(16)
    for run in p.runs:
        run.bold=True;prop(run._r.get_or_add_rPr(),'bCs',val=1)
    p=paragraph(doc,c.arabic,rtl=True,style='Title',size=22)
    p.paragraph_format.space_after=Pt(4)
    p=paragraph(doc,c.english,style='Subtitle',size=14)
    p.paragraph_format.space_after=Pt(18)
    for key in ('certificate.statement.ar','certificate.statement.en'):
        p=paragraph(doc,'{{'+key+'}}',size=11)
        p.paragraph_format.line_spacing=1.3;p.paragraph_format.space_after=Pt(14)
    t=table(doc,[41,72,61]);header=t.add_row()
    for cell,label in zip(header.cells,(('الدور','Role'),('الاسم','Name'),('الوحدة التنظيمية','Organizational unit'))):
        cell._tc.remove(cell.paragraphs[0]._p)
        prop(cell._tc.get_or_add_tcPr(),'shd',fill='EEF7FD')
        paragraph(cell,label[0],rtl=True,size=10);paragraph(cell,label[1],size=9)
    row=t.add_row()
    for i,(cell,key) in enumerate(zip(row.cells,('participant.role','participant.name','participant.org_unit'))):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell._tc.remove(cell.paragraphs[0]._p)
        if i==0:paragraph(cell,'{{#participants}}',size=1)
        paragraph(cell,'{{'+key+'}}',size=10)
        if i==2:paragraph(cell,'{{/participants}}',size=1)
    p=paragraph(doc,'تاريخ إصدار الشهادة / Certificate generation date',rtl=True,size=9)
    p.paragraph_format.space_before=Pt(18)
    paragraph(doc,'{{generated_at}}',size=10)
    doc.core_properties.title='Digital Destruction Certificate';doc.core_properties.subject=c.arabic
    package=BytesIO();doc.save(package)
    with ZipFile(package) as source,ZipFile(directory/c.filename,'w',ZIP_DEFLATED) as output:
        for name in source.namelist():
            info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED
            output.writestr(info,source.read(name))
    print(directory/c.filename)

if __name__=='__main__':main()
